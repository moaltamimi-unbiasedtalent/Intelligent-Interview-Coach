"""P10C-F2A: US official source qualification (SOC 2018 structure and BLS OEWS May 2025 national estimates).

Pure, offline transformations plus two writers that only talk to the graph and the Compensation V2 store they are given (an ISOLATED database chosen by the caller):

* ``parse_soc_structure`` / ``parse_soc_definitions`` read the two official 2018 SOC workbooks; the hierarchy is the PUBLISHED layout (never inferred from titles) and is
  cross-checked against the code prefixes. ``populate_soc_graph`` registers every SOC code as a canonical node and every published parent link as a source-provided ``broader``
  relationship (one transaction, idempotent).
* ``parse_oews_national`` reads the official ``national_M2025_dl.xlsx``; ``oews_v2_records`` builds ``CompensationRecordV2`` rows for every PUBLISHED measurement. Policy:
  - A: BLS-specific aggregate codes (not in the 2018 SOC structure) are QUARANTINED, never SOC nodes: ``CompensationRecordV2`` and the taxonomy vocabulary have no way to carry a
    non-canonical code without misrepresenting it.
  - B: a SOC broad occupation that OEWS lists twice (broad and detailed) with identical values is ONE observation; both source presentations are kept in provenance.
  - C: the BLS relative standard errors travel with every record; no cutoff is invented and no record is qualified here.
  - ``*`` (not available) stays ABSENT: no zero, no interpolation, no hourly-from-annual or annual-from-hourly inference. Published annual wages are never recomputed.
No download, no network, no candidate-facing import."""

from __future__ import annotations

import hashlib
import json
import os
import re
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from src.knowledge_foundation.compensation.record import CompensationRecordV2

__all__ = ["SocNode", "OewsRow", "OEWS_SPECIFIC_CODES", "locate_staging", "verify_staged_files", "parse_soc_structure", "parse_soc_definitions", "check_soc_files", "populate_soc_graph",
           "parse_oews_national", "oews_v2_records", "TRANSFORMATION", "SOC_SOURCE", "OEWS_SOURCE", "SOC_EDITION", "OEWS_EDITION"]

TRANSFORMATION = "p10c-f2a-oews-national-1"
SOC_SOURCE, OEWS_SOURCE = "us_soc_2018", "bls_oews"
SOC_EDITION, OEWS_EDITION = "2018", "May 2025"
OEWS_SPECIFIC_CODES = frozenset({"21-1018", "25-2052", "25-9045", "51-2028", "53-1047"})     # owner policy A (2026-10-18); re-derived from the data by ``oews_v2_records``
PUBLICATION_DATE = datetime(2026, 5, 15, tzinfo=timezone.utc)                                    # news release date, verified on bls.gov
_SOC_CODE = re.compile(r"^\d{2}-\d{4}$")
_LEVELS = ("major", "minor", "broad", "detailed")
_ANNUAL = (("A_MEAN", "mean", 0), ("A_PCT10", "percentile", 10), ("A_PCT25", "percentile", 25), ("A_MEDIAN", "median", 0), ("A_PCT75", "percentile", 75), ("A_PCT90", "percentile", 90))
_HOURLY = (("H_MEAN", "mean", 0), ("H_PCT10", "percentile", 10), ("H_PCT25", "percentile", 25), ("H_MEDIAN", "median", 0), ("H_PCT75", "percentile", 75), ("H_PCT90", "percentile", 90))


@dataclass(frozen=True)
class SocNode:
    code: str
    title: str
    level: str
    parent: str | None
    definition: str = ""
    row: int = 0


@dataclass(frozen=True)
class OewsRow:
    row: int
    values: dict[str, Any]


# ------------------------------------------------------------------------------------------------------------------------------------------ staging and hashes
def locate_staging(manifest_path: Path | None = None) -> Path | None:
    """The approved staging directory: ``P10CF2A_STAGING_DIR``, else the sibling ``ask4mo-source-staging/p10c-f2a/us`` of the repository, else the directory the manifest records.
    Never assumes one absolute path exists."""
    env = os.environ.get("P10CF2A_STAGING_DIR")
    cands = [Path(env)] if env else []
    root = Path(__file__).resolve().parents[2]
    cands.append(root.parent / "ask4mo-source-staging" / "p10c-f2a" / "us")
    if manifest_path and Path(manifest_path).exists():
        m = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
        if m.get("staging_root"):
            cands.append(Path(m["staging_root"]))
    return next((c for c in cands if (c / "raw").is_dir()), None)


def verify_staged_files(manifest: dict[str, Any], staging: Path) -> dict[str, str]:
    """``{file_id: OK | MISSING | HASH_MISMATCH | SIZE_MISMATCH}`` against the committed manifest. Nothing is read beyond the four raw files."""
    out: dict[str, str] = {}
    for f in manifest["files"]:
        p = staging / "raw" / f["original_filename"]
        if not p.exists():
            out[f["file_id"]] = "MISSING"
        elif p.stat().st_size != f["size_bytes"]:
            out[f["file_id"]] = "SIZE_MISMATCH"
        elif hashlib.sha256(p.read_bytes()).hexdigest() != f["sha256"]:
            out[f["file_id"]] = "HASH_MISMATCH"
        else:
            out[f["file_id"]] = "OK"
    return out


# ------------------------------------------------------------------------------------------------------------------------------------------------- SOC 2018
def _sheet_rows(path: Path, sheet: str | None = None) -> list[tuple]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        return [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def parse_soc_structure(path: Path) -> list[SocNode]:
    """The published 2018 SOC structure. A row carries its code in the column of its level (major, minor, broad, detailed); a node's parent is the most recent node of the next level up
    in the published order. Rows that are not codes (title block, blanks) are skipped; a code that does not match xx-xxxx is an error, never repaired."""
    rows = _sheet_rows(path)
    start = next(i for i, r in enumerate(rows) if r[:4] == ("Major Group", "Minor Group", "Broad Group", "Detailed Occupation")) + 1
    nodes: list[SocNode] = []
    last: dict[str, str] = {}
    for n, r in enumerate(rows[start:], start=start + 1):
        cols = [(lvl, r[i]) for i, lvl in enumerate(_LEVELS) if i < len(r) and r[i]]
        if not cols:
            continue
        if len(cols) != 1:
            raise ValueError(f"row {n}: expected exactly one code")
        level, code = cols[0]
        code = str(code).strip()
        if not _SOC_CODE.match(code):
            raise ValueError(f"row {n}: invalid SOC code")
        parent = None if level == "major" else last.get(_LEVELS[_LEVELS.index(level) - 1])
        if level != "major" and parent is None:
            raise ValueError(f"row {n}: no published parent")
        last[level] = code
        for deeper in _LEVELS[_LEVELS.index(level) + 1:]:
            last.pop(deeper, None)
        nodes.append(SocNode(code=code, title=str(r[4] or "").strip(), level=level, parent=parent, row=n))
    return nodes


def parse_soc_definitions(path: Path) -> dict[str, dict[str, str]]:
    """``{code: {level, title, definition}}`` from the official definitions workbook."""
    rows = _sheet_rows(path)
    start = next(i for i, r in enumerate(rows) if r[:4] == ("SOC Group", "SOC Code", "SOC Title", "SOC Definition")) + 1
    out: dict[str, dict[str, str]] = {}
    for r in rows[start:]:
        if r[1]:
            out[str(r[1]).strip()] = {"level": str(r[0] or "").strip().lower(), "title": str(r[2] or "").strip(), "definition": str(r[3] or "").strip()}
    return out


def check_soc_files(structure: list[SocNode], definitions: dict[str, dict[str, str]]) -> dict[str, Any]:
    """Integrity checks between the two official files and the code prefixes. Reports, never repairs."""
    by_code = {n.code: n for n in structure}
    levels = Counter(n.level for n in structure)
    prefix_mismatch = []
    for n in structure:
        if n.parent and n.level != "major":
            span = {"minor": 2, "broad": 4, "detailed": 6}[n.level]           # characters of the code a child shares with its published parent
            if n.parent[:span] != n.code[:span]:
                prefix_mismatch.append(n.code)
    return {"codes": len(structure), "distinct_codes": len(by_code), "duplicate_codes": len(structure) - len(by_code), "levels": dict(levels),
            "definition_codes": len(definitions), "code_sets_equal": set(by_code) == set(definitions),
            "level_disagreements": sorted(c for c, d in definitions.items() if c in by_code and d["level"] != by_code[c].level)[:20],
            "title_disagreements": sum(1 for c, d in definitions.items() if c in by_code and d["title"] != by_code[c].title),
            "parent_prefix_mismatches": prefix_mismatch[:20], "parent_prefix_mismatch_count": len(prefix_mismatch)}


def populate_soc_graph(graph, structure: list[SocNode], definitions: dict[str, dict[str, str]], *, source_ref: str = SOC_SOURCE, edition: str = SOC_EDITION,
                       file_hashes: dict[str, str] | None = None) -> dict[str, int]:
    """Register every SOC code (all four levels, military included) and every published parent link, in ONE transaction. Idempotent: an existing node or identical relationship is left
    unchanged. Needs the ``soc`` taxonomy edition and the registered source to exist."""
    from src.knowledge_foundation.graph import occupation_key_for

    before = graph.counts()
    hashes = file_hashes or {}

    def run(s):
        for n in structure:
            d = definitions.get(n.code, {})
            graph.add_occupation(n.title, [{"scheme": "soc", "edition": edition, "code": n.code, "native_label": n.title}], source_ref=source_ref,
                                 provenance={"soc_level": n.level, "parent_code": n.parent, "definition": d.get("definition", ""), "structure_row": n.row,
                                             "structure_sha256": hashes.get("structure"), "definitions_sha256": hashes.get("definitions"), "transformation": "p10c-f2a-soc-1"}, _session=s)
        for n in structure:
            if n.parent:
                graph.add_relationship(occupation_key_for("soc", edition, n.code), occupation_key_for("soc", edition, n.parent), relationship_type="broader", basis="source_provided",
                                       mapping_status="verified", from_taxonomy=("soc", edition), to_taxonomy=("soc", edition), source_ref=source_ref,
                                       limitations="Published 2018 SOC hierarchy (parent group); a broader group, not an equivalent occupation.", qualification_status="engineering_checked", _session=s)

    graph.bulk(run)
    after = graph.counts()
    return {"nodes_total": after["occupations"], "identifiers_total": after["identifiers"], "relationships_total": after["relationships"],
            "nodes_added": after["occupations"] - before["occupations"], "relationships_added": after["relationships"] - before["relationships"]}


# ------------------------------------------------------------------------------------------------------------------------------------------------ OEWS May 2025
def parse_oews_national(path: Path) -> tuple[list[OewsRow], dict[str, Any]]:
    """The national data sheet plus the BLS-authored field descriptions. Header names are upper-case in the data; every column the dictionary defines must be present."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        data = [tuple(r) for r in wb["national_M2025_dl"].iter_rows(values_only=True)]
        fields: dict[str, str] = {}
        inside = False
        for r in wb["Field Descriptions"].iter_rows(values_only=True):
            if r and r[0] == "Field":
                inside = True
            elif r and r[0] == "Notes:":
                inside = False
            elif inside and r and r[0] and r[1]:
                fields[str(r[0]).strip().lower()] = str(r[1]).strip()
    finally:
        wb.close()
    header = [str(h) for h in data[0]]
    rows = [OewsRow(row=i, values=dict(zip(header, r))) for i, r in enumerate(data[1:], start=2)]
    return rows, {"header": header, "field_descriptions": fields}


def _amount(value: Any) -> tuple[Decimal | None, str]:
    """(amount, state): ``reported``, ``not_available`` (the official ``*``), ``employment_not_available`` (``**``), ``top_coded`` (``#``), ``invalid``. No other string is accepted."""
    if value is None:
        return None, "invalid"
    if isinstance(value, str):
        t = value.strip()
        return None, {"*": "not_available", "**": "employment_not_available", "#": "top_coded"}.get(t, "invalid")
    if isinstance(value, bool):
        return None, "invalid"
    d = Decimal(repr(value)) if isinstance(value, float) else Decimal(str(value))
    return (d, "reported") if d.is_finite() and d > 0 else (None, "invalid")


def oews_v2_records(rows: Iterable[OewsRow], soc_levels: dict[str, str], *, archive_sha256: str, member: str = "national_M2025_dl.xlsx", manifest_file_id: str = "us_oews_may2025_national",
                    ) -> tuple[list[CompensationRecordV2], dict[str, Any]]:
    """Build V2 records from the official rows. ``soc_levels`` is ``{code: level}`` from the SOC structure (the authority for what a code IS). Returns (records, report); the report holds
    counts, the quarantined codes and every absent measurement, and never any salary value."""
    rep: dict[str, Any] = {"rows": 0, "total_rows": 0, "quarantined_rows": 0, "quarantined_codes": {}, "duplicate_presentations": {}, "occupations": 0, "records": 0, "absent_measurements": Counter(),
                           "absent_by_flag": Counter(), "special_markers": Counter(), "rejected": Counter(), "level_counts": Counter(), "annual_only_occupations": 0, "hourly_only_occupations": 0}
    by_code: dict[str, list[OewsRow]] = {}
    for r in rows:
        rep["rows"] += 1
        group = str(r.values.get("O_GROUP") or "").strip().lower()
        code = str(r.values.get("OCC_CODE") or "").strip()
        if group == "total":
            rep["total_rows"] += 1
            continue
        if not _SOC_CODE.match(code):
            rep["rejected"]["invalid_occupation_code"] += 1
            continue
        if code not in soc_levels:
            rep["quarantined_rows"] += 1
            rep["quarantined_codes"][code] = {"title": r.values.get("OCC_TITLE"), "o_group": group, "row": r.row, "reason": "bls_specific_aggregate_not_in_soc_2018"}
            continue
        by_code.setdefault(code, []).append(r)
    out: list[CompensationRecordV2] = []
    for code, group_rows in sorted(by_code.items()):
        first = group_rows[0]
        if len(group_rows) > 1:
            same = all({k: v for k, v in g.values.items() if k != "O_GROUP"} == {k: v for k, v in first.values.items() if k != "O_GROUP"} for g in group_rows[1:])
            if not same:
                rep["rejected"]["conflicting_duplicate_presentations"] += 1
                continue
            rep["duplicate_presentations"][code] = [str(g.values.get("O_GROUP")) for g in group_rows]
        soc_level = soc_levels[code]
        for g in group_rows:
            og = str(g.values.get("O_GROUP") or "").strip().lower()
            if og != soc_level and not (og == "detailed" and len(group_rows) > 1):
                rep["rejected"]["o_group_disagrees_with_soc_structure"] += 1
        level = "detailed" if soc_level == "detailed" else "broad"
        v = first.values
        rep["occupations"] += 1
        rep["level_counts"][level] += 1
        annual_only, hourly_only = str(v.get("ANNUAL") or "").strip().upper() == "TRUE", str(v.get("HOURLY") or "").strip().upper() == "TRUE"
        rep["annual_only_occupations"] += annual_only
        rep["hourly_only_occupations"] += hourly_only
        meta = {"transformation": TRANSFORMATION, "manifest_file_id": manifest_file_id, "archive": "oesm25nat.zip", "archive_sha256": archive_sha256, "archive_member": member,
                "sheet": "national_M2025_dl", "source_rows": [g.row for g in group_rows], "source_o_groups": [str(g.values.get("O_GROUP")) for g in group_rows], "occ_title": v.get("OCC_TITLE"),
                "tot_emp": v.get("TOT_EMP"), "emp_prse": v.get("EMP_PRSE"), "mean_prse": v.get("MEAN_PRSE"), "annual_only": annual_only, "hourly_only": hourly_only,
                "soc_level": soc_level, "naics": v.get("NAICS"), "own_code": v.get("OWN_CODE"), "area": v.get("AREA")}
        for period, spec in (("annual", _ANNUAL), ("hourly", _HOURLY)):
            amounts = {col: _amount(v.get(col)) for col, _, _ in spec}
            for col, (_, state) in amounts.items():
                if state not in ("reported",):
                    key = f"{col}:{state}"
                    rep["absent_measurements"][key] += 1
                    rep["absent_by_flag"][("annual_only" if annual_only else "hourly_only" if hourly_only else "no_flag") + ":" + period] += 1
                    if state != "not_available":
                        rep["special_markers"][key] += 1
            for col, measure, pct in spec:
                val, state = amounts[col]
                if state != "reported":
                    continue                                                       # absent stays absent: no record, no zero, no inference
                lower = upper = None
                if measure == "median":
                    lo, hi = amounts[f"{col[0]}_PCT10"], amounts[f"{col[0]}_PCT90"]
                    if lo[1] == "reported" and hi[1] == "reported" and lo[0] <= hi[0]:
                        lower, upper = lo[0], hi[0]
                try:
                    rec = CompensationRecordV2(
                        taxonomy_scheme="soc", taxonomy_edition=SOC_EDITION, native_code=code, occupation_level=level, country="US", geography_level="national", currency="USD",
                        pay_period=period, pay_basis="gross", measure=measure, percentile=pct, earnings_concept="wage_statistic", population_scope="employees", evidence_method="survey",
                        value_status="reported", value=val, lower_value=lower, upper_value=upper, bounds_semantic="percentile_10_90" if lower is not None else "unknown",
                        source_ref=OEWS_SOURCE, source_edition=OEWS_EDITION, reference_year=2025, reference_period="May 2025", publication_date=PUBLICATION_DATE,
                        limitations="BLS OEWS May 2025 national, cross-industry; straight-time gross wage excluding premium pay; employees only (self-employed excluded). Published value, not recomputed.",
                        provenance={**meta, "source_field": col, "annual_basis": (None if period == "hourly" else "annual_only_occupation_reported_annual" if annual_only
                                                                                    else "bls_published_annual_may_be_derived_from_hourly_at_2080_hours")})
                except ValueError:
                    rep["rejected"]["schema_invalid"] += 1
                    continue
                out.append(rec)
    rep["records"] = len(out)
    for k in ("absent_measurements", "absent_by_flag", "special_markers", "rejected", "level_counts"):
        rep[k] = dict(rep[k])
    return out, rep
