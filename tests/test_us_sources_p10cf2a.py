"""P10C-F2A: US SOC 2018 and BLS OEWS ingestion mechanics. FULLY OFFLINE: tiny SYNTHETIC workbooks in the official layouts (invented codes beyond structure, invented amounts), throwaway SQLite
databases. Nothing here is official data. Tests that read the real staged BLS files are skipped when the staging directory is absent."""

from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from src.knowledge_foundation import us_sources as U
from src.knowledge_foundation.aliases import AliasStore
from src.knowledge_foundation.compensation.resolver import CompensationOutcome, CompensationRequest, CompensationResolver
from src.knowledge_foundation.compensation.store import CompensationSourceNotApproved, CompensationStoreV2
from src.knowledge_foundation.graph import OccupationGraph
from src.knowledge_foundation.occupation_resolver import MultilingualRequest, OccupationResolver
from src.knowledge_foundation.registry import SourceRegistry
from src.knowledge_foundation.taxonomy import TaxonomyCatalogue
from src.persistence import init_db, make_engine, make_session_factory

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "docs/capstone/p10/p10c_f2a/acquisition_manifest.json"
HDR = ["AREA", "AREA_TITLE", "AREA_TYPE", "PRIM_STATE", "NAICS", "NAICS_TITLE", "I_GROUP", "OWN_CODE", "OCC_CODE", "OCC_TITLE", "O_GROUP", "TOT_EMP", "EMP_PRSE", "JOBS_1000", "LOC_QUOTIENT",
       "PCT_TOTAL", "PCT_RPT", "H_MEAN", "A_MEAN", "MEAN_PRSE", "H_PCT10", "H_PCT25", "H_MEDIAN", "H_PCT75", "H_PCT90", "A_PCT10", "A_PCT25", "A_MEDIAN", "A_PCT75", "A_PCT90", "ANNUAL", "HOURLY"]


def _structure(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2018 Structure"
    ws.append(["Synthetic header"])
    ws.append([])
    ws.append(["Major Group", "Minor Group", "Broad Group", "Detailed Occupation", None])
    rows = [("11-0000", None, None, None, "Syn Major"), (None, "11-1000", None, None, "Syn Minor"), (None, None, "11-1010", None, "Syn Broad"), (None, None, None, "11-1011", "Syn Detailed"),
            (None, None, None, "11-1012", "Syn Detailed Two"), (None, None, "11-1020", None, "Syn Broad Two"), (None, None, None, "11-1021", "Syn Detailed Three"),
            ("55-0000", None, None, None, "Syn Military"), (None, "55-1000", None, None, "Syn Military Minor"), (None, None, "55-1010", None, "Syn Military Broad"), (None, None, None, "55-1011", "Syn Military Detailed")]
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def _definitions(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2018 SOC"
    ws.append(["Synthetic"])
    ws.append(["SOC Group", "SOC Code", "SOC Title", "SOC Definition"])
    for g, c, t in (("Major", "11-0000", "Syn Major"), ("Minor", "11-1000", "Syn Minor"), ("Broad", "11-1010", "Syn Broad"), ("Detailed", "11-1011", "Syn Detailed"), ("Detailed", "11-1012", "Syn Detailed Two"),
                    ("Broad", "11-1020", "Syn Broad Two"), ("Detailed", "11-1021", "Syn Detailed Three"), ("Major", "55-0000", "Syn Military"), ("Minor", "55-1000", "Syn Military Minor"),
                    ("Broad", "55-1010", "Syn Military Broad"), ("Detailed", "55-1011", "Syn Military Detailed")):
        ws.append([g, c, t, "synthetic definition" if g == "Detailed" else ""])
    wb.save(path)


def _row(code, group, title="T", **kw):
    base = {"AREA": 99, "AREA_TITLE": "U.S.", "AREA_TYPE": 1, "PRIM_STATE": "US", "NAICS": "000000", "NAICS_TITLE": "Cross-industry", "I_GROUP": "cross-industry", "OWN_CODE": 1235, "OCC_CODE": code,
            "OCC_TITLE": title, "O_GROUP": group, "TOT_EMP": 1000, "EMP_PRSE": 1.5, "H_MEAN": 20.5, "A_MEAN": 42640, "MEAN_PRSE": 0.7, "H_PCT10": 12.1, "H_PCT25": 15.25, "H_MEDIAN": 19.9, "H_PCT75": 25.0,
            "H_PCT90": 31.5, "A_PCT10": 25168, "A_PCT25": 31720, "A_MEDIAN": 41392, "A_PCT75": 52000, "A_PCT90": 65520, "ANNUAL": None, "HOURLY": None}
    base.update(kw)
    return [base.get(h) for h in HDR]


def _oews(path: Path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "national_M2025_dl"
    ws.append(HDR)
    for r in rows:
        ws.append(r)
    fd = wb.create_sheet("Field Descriptions")
    for r in (["May 2025 OEWS Estimates"], ["Field", "Field Description"], ["occ_code", "The 6-digit SOC code"], ["Notes:"], ["* = a wage estimate is not available"]):
        fd.append(r)
    wb.save(path)


@pytest.fixture()
def files(tmp_path):
    s, d, o = tmp_path / "s.xlsx", tmp_path / "d.xlsx", tmp_path / "o.xlsx"
    _structure(s)
    _definitions(d)
    rows = [_row("00-0000", "total"), _row("11-0000", "major"), _row("11-1000", "minor"), _row("11-1010", "broad"), _row("11-1010", "detailed"), _row("11-1011", "detailed"),
            _row("11-1012", "detailed", ANNUAL="TRUE", H_MEAN="*", H_PCT10="*", H_PCT25="*", H_MEDIAN="*", H_PCT75="*", H_PCT90="*"),
            _row("11-1021", "detailed", HOURLY="TRUE", A_MEAN="*", A_PCT10="*", A_PCT25="*", A_MEDIAN="*", A_PCT75="*", A_PCT90="*"),
            _row("21-1018", "detailed", "Syn BLS aggregate")]
    _oews(o, rows)
    return s, d, o


@pytest.fixture()
def sf(tmp_path):
    engine = make_engine(f"sqlite:///{tmp_path / 'c.db'}")
    init_db(engine)
    return make_session_factory(engine)


def _world(sf, files):
    s, d, o = files
    tax, reg, g = TaxonomyCatalogue(sf), SourceRegistry(sf), OccupationGraph(sf)
    tax.seed()
    tax.ensure("soc", "2018", status="current", acquisition_status="available_locally", code_pattern="[0-9]{2}-[0-9]{4}")
    for sid in (U.SOC_SOURCE, U.OEWS_SOURCE):
        reg.register(sid, origin="authorised_manual", publisher="Synthetic official publisher", dataset_name="Synthetic dataset", dataset_version="E", source_url="https://example.invalid/",
                     precedence_category="national_official_occupation_statistics", validation_status="unvalidated")
    structure, defs = U.parse_soc_structure(s), U.parse_soc_definitions(d)
    return tax, reg, g, structure, defs


def _approve(reg, sid):
    reg.transition(sid, "licensing_pending")
    reg.set_licence(sid, licence_id="syn", licence_status="verified")
    reg.transition(sid, "licensing_verified")
    reg.transition(sid, "approved_for_ingestion")


# ---------------------------------------------------------------------------------------------------------------------------------------------------------- SOC
def test_structure_hierarchy_is_the_published_layout(files):
    st = U.parse_soc_structure(files[0])
    by = {n.code: n for n in st}
    assert [n.level for n in st if n.code.startswith("11")][:4] == ["major", "minor", "broad", "detailed"]
    assert by["11-1011"].parent == "11-1010" and by["11-1021"].parent == "11-1020" and by["11-1010"].parent == "11-1000" and by["11-0000"].parent is None
    assert by["55-1011"].level == "detailed"                                         # military occupations are kept
    chk = U.check_soc_files(st, U.parse_soc_definitions(files[1]))
    assert chk["code_sets_equal"] and chk["duplicate_codes"] == 0 and chk["level_disagreements"] == []


def test_invalid_soc_code_is_an_error_not_repaired(tmp_path):
    p = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Major Group", "Minor Group", "Broad Group", "Detailed Occupation", None])
    ws.append(["11-000", None, None, None, "bad"])
    wb.save(p)
    with pytest.raises(ValueError):
        U.parse_soc_structure(p)


def test_soc_graph_population_is_idempotent_and_hierarchical(sf, files):
    _, reg, g, st, defs = _world(sf, files)
    first = U.populate_soc_graph(g, st, defs)
    again = U.populate_soc_graph(g, st, defs)
    assert first["nodes_total"] == len(st) == 11 and first["relationships_total"] == 9
    assert again["nodes_added"] == 0 and again["relationships_added"] == 0 and again["nodes_total"] == 11
    node = g.find_by_identifier("soc", "11-1011", "2018")[0]
    rel = g.relationships(node["occupation_key"])
    assert len(rel) == 1 and rel[0]["relationship_type"] == "broader" and rel[0]["basis"] == "source_provided" and rel[0]["mapping_status"] == "verified"
    assert g.find_by_identifier("soc", "11-1011", "2010") == []                       # another edition is another namespace
    assert g.identifiers(node["occupation_key"])[0]["native_label"] == "Syn Detailed"


# --------------------------------------------------------------------------------------------------------------------------------------------------------- OEWS
def test_oews_transform_policies(files):
    st = U.parse_soc_structure(files[0])
    rows, meta = U.parse_oews_national(files[2])
    assert meta["field_descriptions"]["occ_code"].startswith("The 6-digit")
    recs, rep = U.oews_v2_records(rows, {n.code: n.level for n in st}, archive_sha256="a" * 64)
    assert rep["total_rows"] == 1                                                  # the all-occupations total is not an occupation
    assert set(rep["quarantined_codes"]) == {"21-1018"}                            # policy A: not canonical, not imported, not mapped
    assert rep["duplicate_presentations"] == {"11-1010": ["broad", "detailed"]}   # policy B: one observation, both presentations recorded
    b = [r for r in recs if r.native_code == "11-1010" and r.pay_period == "annual" and r.measure == "median"]
    assert len(b) == 1 and b[0].occupation_level == "broad" and b[0].provenance["source_o_groups"] == ["broad", "detailed"]
    assert rep["annual_only_occupations"] == 1 and rep["hourly_only_occupations"] == 1
    annual_only = [r for r in recs if r.native_code == "11-1012"]
    assert {r.pay_period for r in annual_only} == {"annual"}                      # '*' stays absent: no hourly record, no inference from annual
    hourly_only = [r for r in recs if r.native_code == "11-1021"]
    assert {r.pay_period for r in hourly_only} == {"hourly"}
    assert all(r.value is not None and r.value > 0 for r in recs)                  # no zero, no suppressed rows
    assert rep["absent_measurements"]["H_MEDIAN:not_available"] == 1 and rep["absent_measurements"]["A_MEDIAN:not_available"] == 1


def test_values_are_exact_distinct_and_carry_reliability(files):
    st = U.parse_soc_structure(files[0])
    rows, _ = U.parse_oews_national(files[2])
    recs, _ = U.oews_v2_records(rows, {n.code: n.level for n in st}, archive_sha256="b" * 64)
    d = {(r.native_code, r.pay_period, r.measure, r.percentile): r for r in recs if r.native_code == "11-1011"}
    assert d[("11-1011", "hourly", "median", 0)].value == Decimal("19.9") and d[("11-1011", "hourly", "mean", 0)].value == Decimal("20.5")
    assert d[("11-1011", "annual", "median", 0)].value == Decimal("41392")           # published annual value, never hourly x 2080
    assert d[("11-1011", "hourly", "percentile", 25)].value == Decimal("15.25")
    med = d[("11-1011", "annual", "median", 0)]
    assert med.lower_value == Decimal("25168") and med.upper_value == Decimal("65520")
    assert med.pay_basis == "gross" and med.earnings_concept == "wage_statistic" and med.population_scope == "employees" and med.evidence_method == "survey"
    assert med.provenance["emp_prse"] == 1.5 and med.provenance["mean_prse"] == 0.7 and med.provenance["tot_emp"] == 1000      # uncertainty preserved, no cutoff applied
    assert med.provenance["archive_sha256"] == "b" * 64 and med.source_ref == "bls_oews" and med.reference_period == "May 2025"
    assert med.source_edition == "May 2025" and med.taxonomy_edition == "2018"


def test_special_markers_other_than_the_documented_ones_are_not_accepted():
    assert U._amount("*") == (None, "not_available") and U._amount("**")[1] == "employment_not_available" and U._amount("#")[1] == "top_coded"
    assert U._amount("n/a")[1] == "invalid" and U._amount("") [1] == "invalid" and U._amount(0)[1] == "invalid" and U._amount(None)[1] == "invalid"


def test_conflicting_duplicate_presentations_are_rejected(files, tmp_path):
    st = U.parse_soc_structure(files[0])
    o = tmp_path / "dup.xlsx"
    _oews(o, [_row("11-1010", "broad"), _row("11-1010", "detailed", A_MEDIAN=1)])
    rows, _ = U.parse_oews_national(o)
    recs, rep = U.oews_v2_records(rows, {n.code: n.level for n in st}, archive_sha256="c" * 64)
    assert recs == [] and rep["rejected"]["conflicting_duplicate_presentations"] == 1


# ------------------------------------------------------------------------------------------------------------- governed import, lifecycle, retrieval, revocation
def test_import_is_registry_gated_idempotent_and_never_operational(sf, files):
    tax, reg, g, st, defs = _world(sf, files)
    U.populate_soc_graph(g, st, defs)
    rows, _ = U.parse_oews_national(files[2])
    recs, _ = U.oews_v2_records(rows, {n.code: n.level for n in st}, archive_sha256="d" * 64)
    store = CompensationStoreV2(sf, reg, g)
    with pytest.raises(CompensationSourceNotApproved):
        store.import_records(recs, source_ref=U.OEWS_SOURCE, source_edition="May 2025", mode="production")          # registered but not approved: refused
    _approve(reg, U.OEWS_SOURCE)
    first = store.import_records(recs, source_ref=U.OEWS_SOURCE, source_edition="May 2025", mode="production", transformation=U.TRANSFORMATION)
    again = store.import_records(recs, source_ref=U.OEWS_SOURCE, source_edition="May 2025", mode="production", transformation=U.TRANSFORMATION)
    assert first.added == len(recs) and first.rejected == 0 and again.already_imported and again.added == 0 and store.counts()["records"] == len(recs)
    assert not reg.can_serve_as_evidence(U.OEWS_SOURCE)                                                              # ingestion approval is not runtime approval
    occ = OccupationResolver(g, AliasStore(sf), registry=reg)
    res = CompensationResolver(store, occ, g, reg)

    def ask(code, **kw):
        return res.resolve(CompensationRequest(occupation=MultilingualRequest(taxonomy_scheme="soc", taxonomy_code=code, taxonomy_edition="2018"), country="US", require_operational=False, **kw))

    exact = ask("11-1011", pay_period="annual", measure="median")
    assert exact.outcome is CompensationOutcome.EVIDENCE and exact.evidence_class == "C1" and exact.value == "41392" and exact.citation is not None
    assert ask("11-1010", pay_period="annual", measure="median").evidence_class == "C2"                              # broad occupation stays a broad (C2) claim
    assert ask("21-1018", pay_period="annual", measure="median").outcome is CompensationOutcome.NOT_FOUND            # non-canonical aggregate: no identity, no evidence
    assert ask("99-9999", pay_period="annual", measure="median").outcome is CompensationOutcome.NOT_FOUND
    assert ask("55-1011", pay_period="annual", measure="median").outcome is CompensationOutcome.ABSTAIN             # military: in the graph, no published wage, C4
    operational = res.resolve(CompensationRequest(occupation=MultilingualRequest(taxonomy_scheme="soc", taxonomy_code="11-1011", taxonomy_edition="2018"), country="US", pay_period="annual", measure="median"))
    assert operational.outcome is CompensationOutcome.ABSTAIN and operational.value is None and not operational.candidate_presentable


def test_revocation_blocks_further_ingestion_and_keeps_records(sf, files):
    tax, reg, g, st, defs = _world(sf, files)
    U.populate_soc_graph(g, st, defs)
    rows, _ = U.parse_oews_national(files[2])
    recs, _ = U.oews_v2_records(rows, {n.code: n.level for n in st}, archive_sha256="e" * 64)
    store = CompensationStoreV2(sf, reg, g)
    _approve(reg, U.OEWS_SOURCE)
    store.import_records(recs, source_ref=U.OEWS_SOURCE, source_edition="May 2025", mode="production")
    reg.transition(U.OEWS_SOURCE, "restricted", reason="synthetic revocation")
    assert not reg.can_ingest(U.OEWS_SOURCE) and not reg.can_serve_as_evidence(U.OEWS_SOURCE)
    with pytest.raises(CompensationSourceNotApproved):
        store.import_records(recs, source_ref=U.OEWS_SOURCE, source_edition="May 2025", mode="production")
    assert store.counts()["records"] == len(recs)                                                                    # nothing silently deleted
    assert [e["to"] for e in reg.events(U.OEWS_SOURCE)][-1] == "restricted"


# ----------------------------------------------------------------------------------------------------------------------------------------- staged official files
def _staging():
    return U.locate_staging(MANIFEST)


@pytest.mark.skipif(_staging() is None, reason="staged BLS files not present in this environment")
def test_staged_files_match_the_committed_manifest():
    m = json.loads(MANIFEST.read_text())
    assert set(U.verify_staged_files(m, _staging()).values()) == {"OK"}


@pytest.mark.skipif(_staging() is None, reason="staged BLS files not present in this environment")
def test_real_soc_and_oews_counts():
    st = _staging()
    structure = U.parse_soc_structure(st / "raw/soc_structure_2018.xlsx")
    chk = U.check_soc_files(structure, U.parse_soc_definitions(st / "raw/soc_2018_definitions.xlsx"))
    assert chk["levels"] == {"major": 23, "minor": 98, "broad": 459, "detailed": 867} and chk["code_sets_equal"] and chk["title_disagreements"] == 0
    rows, meta = U.parse_oews_national(st / "extracted/oesm25nat/oesm25nat/national_M2025_dl.xlsx")
    assert len(rows) == 1401 and len(meta["field_descriptions"]) >= 30
    recs, rep = U.oews_v2_records(rows, {n.code: n.level for n in structure}, archive_sha256="0" * 64)
    assert set(rep["quarantined_codes"]) == U.OEWS_SPECIFIC_CODES and len(rep["duplicate_presentations"]) == 7
    assert rep["occupations"] == 1388 and rep["records"] == 16140 and rep["annual_only_occupations"] == 79 and rep["hourly_only_occupations"] == 7 and rep["rejected"] == {}
    assert rep["special_markers"] == {}                                                                              # only '*' occurs; ** and # do not
